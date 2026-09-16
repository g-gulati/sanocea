from __future__ import annotations

import shutil
from pathlib import Path
import zipfile

import openpyxl
import pymupdf


def generate_reference_merchant_fixtures(target_dir: Path) -> dict[str, Path]:
    target_dir.mkdir(parents=True, exist_ok=True)
    generated: dict[str, Path] = {}

    # 1. supplier_price_list_messy.xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price_List"
    headers = [
        "Parent SKU", "Title", "Product Type", "Category", "Variant SKU",
        "Packaging", "Volume", "Price Formula", "Barcode", "Weight", "Dimensions", "HSN"
    ]
    ws.append(headers)

    # Style: ANCHAL-KACHI-GHANI with 3 pack sizes and merged cells
    ws.append([
        "ANCHAL-KACHI-GHANI", "Anchal Cold-Pressed Kachi Ghani Mustard Oil", "Cooking Oil", "Mustard Oil",
        "ANCHAL-KACHI-1L-PCH", "Pouch", "1000ml", "=130*1.20", "8901234567011", "910g", "15x10x5 cm", "15141910"
    ])
    ws.append([
        None, None, None, None,
        "ANCHAL-KACHI-1L-BTL", "Bottle", "1000ml", "=140*1.20", "8901234567028", "950g", "8x8x28 cm", "15141910"
    ])
    ws.append([
        None, None, None, None,
        "ANCHAL-KACHI-5L-CAN", "Can", "5000ml", "=650*1.20", "8901234567035", "4600g", "20x15x32 cm", "15141910"
    ])

    ws.merge_cells("A2:A4")
    ws.merge_cells("B2:B4")
    ws.merge_cells("C2:C4")
    ws.merge_cells("D2:D4")
    ws.merge_cells("L2:L4")

    # Sheet 2: Hidden sheet containing internal cost and procurement margins
    ws_hidden = wb.create_sheet(title="Internal_Costing")
    ws_hidden.sheet_state = "hidden"
    ws_hidden.append(["Style", "Procurement Mill Cost", "Target Gross Margin", "FSSAI License"])
    ws_hidden.append(["ANCHAL-KACHI-GHANI", "98.50 INR/L", "38.5%", "10014011002233"])

    xlsx_path = target_dir / "supplier_price_list_messy.xlsx"
    wb.save(xlsx_path)

    # Patch cached values for formulas in sheet1.xml
    temp_zip = target_dir / "temp_ref.zip"
    shutil.copyfile(xlsx_path, temp_zip)
    with zipfile.ZipFile(temp_zip, "r") as zin, zipfile.ZipFile(xlsx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = content.decode("utf-8")
                # H2: 156.0, H3: 168.0, H4: 780.0
                text = text.replace('<c r="H2"><f>130*1.20</f><v></v></c>', '<c r="H2"><f>130*1.20</f><v>156.0</v></c>')
                text = text.replace('<c r="H3"><f>140*1.20</f><v></v></c>', '<c r="H3"><f>140*1.20</f><v>168.0</v></c>')
                text = text.replace('<c r="H4"><f>650*1.20</f><v></v></c>', '<c r="H4"><f>650*1.20</f><v>780.0</v></c>')
                content = text.encode("utf-8")
            zout.writestr(item, content)
    temp_zip.unlink(missing_ok=True)
    generated["supplier_price_list"] = xlsx_path

    # 2. product_specifications.pdf (Multi-product PDF with PyMuPDF)
    pdf_doc = pymupdf.open()
    page1 = pdf_doc.new_page(width=612, height=792)
    spec_text_p1 = """# Spec Catalog - Page 1
| SKU | Title | Price | Type | Material | HSN |
|---|---|---|---|---|---|
| ANCHAL-DESI-GHEE | Anchal Vedic A2 Cow Ghee | 1499.00 | Dairy | Pure Butterfat | 04059020 |
"""
    page1.insert_text((40, 50), spec_text_p1, fontsize=9)

    page2 = pdf_doc.new_page(width=612, height=792)
    spec_text_p2 = """# Spec Catalog - Page 2
| SKU | Title | Price | Type | Material | HSN |
|---|---|---|---|---|---|
| ANCHAL-WILD-HONEY | Anchal Himalayan Honey | 650.00 | Honey | Raw Honey | 04090000 |
"""
    page2.insert_text((40, 50), spec_text_p2, fontsize=9)

    pdf_path = target_dir / "product_specifications.pdf"
    pdf_doc.save(pdf_path)
    pdf_doc.close()
    generated["product_specifications"] = pdf_path

    # 3. marketplace_feed_jagged.csv
    # Indian formatting, missing units, style ANCHAL-WOOD-SESAME
    csv_content = """sku,title,price,currency,product_type,weight,dimensions,packaging,hsn
ANCHAL-WOOD-SESAME,Anchal Wood Pressed Black Sesame Oil,"₹1,299.00",INR,Cooking Oil,920g,10x10x26,Glass Bottle,15155091
JAGGED-ROW-EXTRA,Incomplete Extra Row Item,"₹499.00",INR
,Cold Pressed Castor Oil Pure,"₹350.00",INR,Personal Care,200g,5x5x15 cm,Amber Bottle,15153090
"""
    csv_path = target_dir / "marketplace_feed_jagged.csv"
    csv_path.write_text(csv_content, encoding="utf-8-sig")
    generated["marketplace_feed"] = csv_path

    # 4. conflicting_feed.csv
    # Conflicting price for ANCHAL-KACHI-GHANI to trigger CONFLICTED state and operator review
    conflict_csv = """parent_sku,sku,title,price,currency,product_type
ANCHAL-KACHI-GHANI,ANCHAL-KACHI-1L-PCH,Anchal Cold-Pressed Kachi Ghani Mustard Oil,145.00,INR,Cooking Oil
"""
    conflict_path = target_dir / "conflicting_feed.csv"
    conflict_path.write_text(conflict_csv, encoding="utf-8")
    generated["conflicting_feed"] = conflict_path

    # 5. Ballast files to verify fault isolation and quarantine
    corrupt_bin = target_dir / "ballast_corrupt.bin"
    corrupt_bin.write_bytes(b"\x00\xff\xfe\x00CORRUPT_SUPPLIER_PAYLOAD\x00\x00")
    generated["ballast_corrupt"] = corrupt_bin

    ds_store = target_dir / ".DS_Store"
    ds_store.write_bytes(b"\x00\x00\x00\x01Bud1\x00\x00\x10\x00")
    generated["ds_store"] = ds_store

    return generated
