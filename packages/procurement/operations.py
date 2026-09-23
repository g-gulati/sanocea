from __future__ import annotations

import csv
from dataclasses import asdict, dataclass, field
from datetime import timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from sanocea.packages.audit import AuditLedger
from sanocea.packages.connector_sdk import MutationRequest
from sanocea.packages.domain_contract.location import resolve_location_ref
from sanocea.packages.domain_contract.models import (
    Approval,
    ExceptionRecord,
    GoodsReceipt,
    GoodsReceiptLine,
    InboundShipment,
    InboundShipmentLine,
    Inventory,
    Order,
    OrderLine,
    PurchaseOrder,
    PurchaseOrderLine,
    ReplenishmentRecommendation,
    Supplier,
    SupplierAcknowledgement,
    SupplierSku,
    now_utc,
)
from sanocea.packages.exceptions import ExceptionService
from sanocea.packages.policy_engine.engine import Decision

_STRICTNESS = {Decision.ALLOW: 0, Decision.REQUIRE_APPROVAL: 1, Decision.DENY: 2}


def _stricter(a: Decision, b: Decision) -> Decision:
    return a if _STRICTNESS[a] >= _STRICTNESS[b] else b


_REQUIRED_OFFER_FIELDS = ("sku", "supplier_sku", "cost", "currency", "moq", "pack_quantity", "lead_time_days")


@dataclass
class ProcurementCounters:
    """Kept intentionally close to FinanceCounters' shape - same reporting discipline."""

    supplier_offers_ingested: int = 0
    supplier_offer_rows_rejected: int = 0
    replenishment_recommendations: int = 0
    replenishment_recommendations_needed: int = 0
    purchase_orders_created: int = 0
    auto_approved: int = 0
    requires_approval: int = 0
    blocked: int = 0
    submitted: int = 0
    acknowledged_full: int = 0
    acknowledged_partial: int = 0
    acknowledged_rejected: int = 0
    stale_events_ignored: int = 0
    goods_receipts: int = 0
    goods_receipt_lines_matched: int = 0
    goods_receipt_lines_shortage: int = 0
    goods_receipt_lines_excess: int = 0
    goods_receipt_lines_wrong_sku: int = 0
    cost_variance_flagged: int = 0
    corrective_interventions: int = 0
    ai_calls: int = 0

    def as_dict(self) -> dict[str, Any]:
        data = asdict(self)
        po_total = max(self.purchase_orders_created, 1)
        data["auto_approval_rate"] = round(self.auto_approved / po_total, 4)
        return data


class ProcurementService:
    def __init__(self, store, supplier_connector=None) -> None:
        self.store = store
        self.audit = AuditLedger(store)
        self.exceptions = ExceptionService(store)
        self.supplier_connector = supplier_connector

    # ================================================================================================
    # Supplier ingestion - deterministic-first, never invents cost/MOQ/lead time/availability.
    # ================================================================================================

    def upsert_supplier(self, merchant_id: str, name: str, *, default_lead_time_days: int = 7, contact_ref: str | None = None) -> Supplier:
        existing = next((s for s in self.store.list(Supplier, merchant_id) if s.name == name), None)
        if existing:
            return existing
        supplier = Supplier(merchant_id=merchant_id, name=name, default_lead_time_days=default_lead_time_days, contact_ref=contact_ref)
        self.store.put(supplier)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier_ingestion", action="supplier_created", object_type="Supplier", object_id=supplier.id, result="created")
        return supplier

    def ingest_supplier_offer(
        self,
        merchant_id: str,
        supplier_id: str,
        *,
        sku: str,
        supplier_sku: str,
        cost: int,
        currency: str,
        moq: int,
        pack_quantity: int,
        lead_time_days: int,
        preferred: bool = False,
        evidence_ref: str | None = None,
    ) -> SupplierSku:
        """One deterministic supplier offer row. Every field is REQUIRED - this is the structural
        enforcement of "never invent cost/availability/MOQ/lead time/supplier SKU": there is no code
        path that fabricates a default for any of these."""
        offer = SupplierSku(
            merchant_id=merchant_id,
            supplier_id=supplier_id,
            sku=sku,
            supplier_sku=supplier_sku,
            cost=int(cost),
            currency=currency,
            moq=int(moq),
            pack_quantity=int(pack_quantity),
            lead_time_days=int(lead_time_days),
            preferred=preferred,
            evidence_ref=evidence_ref,
        )
        persisted = self.store.put(offer)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier_ingestion", action="supplier_offer_ingested", object_type="SupplierSku", object_id=persisted.id, evidence_ref=evidence_ref, result="observed")
        return persisted

    def ingest_supplier_offer_file(self, merchant_id: str, supplier_id: str, path: Path) -> dict[str, int]:
        """CSV/XLSX -> deterministic parser -> deterministic mapping, mirroring
        packages/product_onboarding/ingestion.py's structured-file pattern. A row missing any required
        field is REJECTED with a ProcurementException, never defaulted."""
        evidence_ref = f"file://{path.resolve()}"
        if path.suffix.lower() == ".csv":
            rows = self._read_csv(path)
        elif path.suffix.lower() in (".xlsx", ".xlsm"):
            rows = self._read_xlsx(path)
        else:
            raise ValueError(f"supplier offer ingestion does not support {path.suffix}")
        ingested = 0
        rejected = 0
        for idx, row in enumerate(rows):
            missing = [field_name for field_name in _REQUIRED_OFFER_FIELDS if not str(row.get(field_name, "")).strip()]
            if missing:
                self.exceptions.create(
                    merchant_id=merchant_id,
                    category="supplier_offer_incomplete_row",
                    message=f"Row {idx} in {path.name} missing required field(s) {missing} - refusing to invent a value",
                    object_id=supplier_id,
                    severity="warning",
                    evidence_ref=evidence_ref,
                )
                rejected += 1
                continue
            try:
                self.ingest_supplier_offer(
                    merchant_id,
                    supplier_id,
                    sku=str(row["sku"]).strip(),
                    supplier_sku=str(row["supplier_sku"]).strip(),
                    cost=int(float(row["cost"])),
                    currency=str(row["currency"]).strip(),
                    moq=int(float(row["moq"])),
                    pack_quantity=int(float(row["pack_quantity"])),
                    lead_time_days=int(float(row["lead_time_days"])),
                    preferred=str(row.get("preferred", "")).strip().lower() in {"1", "true", "yes"},
                    evidence_ref=evidence_ref,
                )
                ingested += 1
            except (TypeError, ValueError) as exc:
                self.exceptions.create(
                    merchant_id=merchant_id, category="supplier_offer_malformed_row",
                    message=f"Row {idx} in {path.name} could not be parsed deterministically: {exc}",
                    object_id=supplier_id, severity="warning", evidence_ref=evidence_ref,
                )
                rejected += 1
        return {"ingested": ingested, "rejected": rejected}

    def _read_csv(self, path: Path) -> list[dict[str, Any]]:
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))

    def _read_xlsx(self, path: Path) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        rows = []
        for raw_row in rows_iter:
            if all(cell is None for cell in raw_row):
                continue
            rows.append({headers[i]: raw_row[i] for i in range(len(headers)) if i < len(raw_row)})
        return rows

    # ================================================================================================
    # Inventory position - explicit on-hand / reserved / confirmed-inbound / available-to-sell split.
    # ================================================================================================

    def inventory_position(self, merchant_id: str, sku: str, location_ref: str = "default") -> dict[str, Any]:
        inv = self._inventory_row(merchant_id, sku, location_ref)
        on_hand = inv.quantity if inv else 0
        reserved = inv.reserved if inv else 0
        confirmed_inbound = inv.confirmed_inbound if inv else 0
        available_to_sell = (inv.available if inv and inv.available is not None else max(on_hand - reserved, 0)) if inv else 0
        return {
            "sku": sku,
            "on_hand": on_hand,
            "reserved": reserved,
            "confirmed_inbound": confirmed_inbound,
            "available_to_sell": available_to_sell,
            # Used for REPLENISHMENT PLANNING ONLY - confirmed inbound reduces how much more to order,
            # but never appears in available_to_sell above (see goods receipt for the only place
            # confirmed inbound converts into sellable stock).
            "inventory_position": on_hand - reserved + confirmed_inbound,
        }

    def _inventory_row(self, merchant_id: str, sku: str, location_ref: str = "default") -> Inventory | None:
        rows = [i for i in self.store.list(Inventory, merchant_id) if i.sku == sku and i.location_ref == location_ref]
        return rows[-1] if rows else None

    # ================================================================================================
    # Replenishment - deterministic, explainable. No ML forecasting.
    # ================================================================================================

    def recent_sales_velocity(self, merchant_id: str, sku: str, lookback_days: int = 30) -> dict[str, Any]:
        cutoff = now_utc() - timedelta(days=lookback_days)
        orders_by_id = {o.id: o for o in self.store.list(Order, merchant_id)}
        units_sold = 0
        for line in self.store.list(OrderLine, merchant_id):
            if line.sku != sku:
                continue
            order = orders_by_id.get(line.order_id)
            if order and order.placed_at and order.placed_at >= cutoff:
                units_sold += line.quantity
        velocity = units_sold / lookback_days if lookback_days else 0.0
        return {"sku": sku, "lookback_days": lookback_days, "units_sold": units_sold, "velocity_per_day": round(velocity, 4)}

    def recommend_replenishment(self, merchant_id: str, sku: str, location_ref: str | None = None) -> ReplenishmentRecommendation:
        """Step 5 (Part C): LOCATION-SPECIFIC replenishment truth. `location_ref` resolves exactly once
        (Step 2's resolve_location_ref - explicit caller value, else merchant-configured default, else
        the legacy "default" literal) and every reading of inventory_position below uses that ONE
        resolved location - a merchant with ample stock at one location must never mask a genuine
        shortage at another (the Surat/Mumbai example in the Step 5 mandate). Sales velocity remains a
        merchant/SKU-wide signal (OrderLine carries no location_ref - fulfilment is not tracked
        per-location on the order side) - this is a deliberate, disclosed scope limit, not an oversight;
        pooling actual on-hand/reserved/ATS across locations is what Part C forbids, not sharing a demand
        signal that has no location dimension to begin with.
        """
        config = self.store.get_config(merchant_id).get("procurement", {})
        per_sku_cfg = config.get("replenishment", {})
        sku_cfg = per_sku_cfg.get(sku, per_sku_cfg.get("default", {}))
        safety_stock = int(sku_cfg.get("safety_stock", 0))
        target_stock_cfg = sku_cfg.get("target_stock")
        static_reorder_point = sku_cfg.get("reorder_point")
        max_stock_cfg = sku_cfg.get("max_stock")
        lookback_days = int(config.get("velocity_lookback_days", 30))

        resolved_location = resolve_location_ref(self.store, merchant_id, explicit=location_ref)
        position = self.inventory_position(merchant_id, sku, resolved_location)
        velocity = self.recent_sales_velocity(merchant_id, sku, lookback_days=lookback_days)
        offer = self.select_supplier_offer(merchant_id, sku)
        lead_time_days = offer.lead_time_days if offer else int(config.get("default_lead_time_days", 7))

        if velocity["units_sold"] > 0:
            reorder_point = safety_stock + round(velocity["velocity_per_day"] * lead_time_days)
            reorder_point_formula = f"safety_stock({safety_stock}) + velocity_per_day({velocity['velocity_per_day']}) * lead_time_days({lead_time_days}) = {reorder_point}"
        elif static_reorder_point is not None:
            reorder_point = int(static_reorder_point)
            reorder_point_formula = f"merchant-configured static reorder_point (no sales in last {lookback_days}d) = {reorder_point}"
        else:
            reorder_point = safety_stock
            reorder_point_formula = f"safety_stock only, no sales history and no configured reorder_point = {reorder_point}"

        needs_reorder = position["inventory_position"] <= reorder_point
        reasoning: dict[str, Any] = {
            "location_ref": resolved_location,
            "inventory_position": position,
            "velocity": velocity,
            "safety_stock": safety_stock,
            "lead_time_days": lead_time_days,
            "lead_time_source": "supplier_offer" if offer else "merchant_default_lead_time_days",
            "reorder_point": reorder_point,
            "reorder_point_formula": reorder_point_formula,
            "needs_reorder": needs_reorder,
        }

        if not needs_reorder:
            reasoning["explanation"] = f"inventory_position({position['inventory_position']}) > reorder_point({reorder_point}) -> no replenishment needed"
            recommended_quantity = 0
        else:
            if target_stock_cfg is not None:
                target_stock = int(target_stock_cfg)
                target_source = "merchant_configured_target_stock"
            else:
                target_stock = reorder_point + round(velocity["velocity_per_day"] * lead_time_days)
                target_source = "reorder_point + one additional lead-time cycle of velocity (no target_stock configured)"
            raw_quantity = max(target_stock - position["inventory_position"], 0)
            reasoning["target_stock"] = target_stock
            reasoning["target_stock_source"] = target_source
            reasoning["raw_quantity_before_moq_pack"] = raw_quantity
            if offer:
                moq_adjusted = max(raw_quantity, offer.moq) if raw_quantity > 0 else 0
                pack = offer.pack_quantity or 1
                packed = ((moq_adjusted + pack - 1) // pack) * pack if moq_adjusted else 0
                reasoning["moq_applied"] = offer.moq
                reasoning["pack_quantity_applied"] = pack
                recommended_quantity = packed
            else:
                recommended_quantity = raw_quantity
                reasoning["moq_applied"] = None
                reasoning["no_active_supplier_offer"] = True
            if max_stock_cfg is not None:
                cap = max(int(max_stock_cfg) - position["inventory_position"], 0)
                reasoning["max_stock_cap"] = int(max_stock_cfg)
                recommended_quantity = min(recommended_quantity, cap)
            reasoning["explanation"] = (
                f"inventory_position({position['inventory_position']}) <= reorder_point({reorder_point}) -> "
                f"target_stock({target_stock}) - position = raw {raw_quantity}, "
                f"MOQ/pack-adjusted to {recommended_quantity}"
            )

        rec = ReplenishmentRecommendation(
            merchant_id=merchant_id, sku=sku, location_ref=resolved_location, recommended_quantity=recommended_quantity,
            chosen_supplier_id=offer.supplier_id if offer else None, reasoning=reasoning,
        )
        self.store.put(rec)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="replenishment", action="replenishment_recommended", object_type="ReplenishmentRecommendation", object_id=rec.id, result=str(recommended_quantity))
        return rec

    # ================================================================================================
    # Deterministic, configurable supplier selection - "no universal best supplier".
    # ================================================================================================

    def select_supplier_offer(self, merchant_id: str, sku: str) -> SupplierSku | None:
        offers = [o for o in self.store.list(SupplierSku, merchant_id) if o.sku == sku and o.status == "active"]
        if not offers:
            return None
        config = self.store.get_config(merchant_id).get("procurement", {})
        priority = config.get("supplier_priority", ["preferred", "cost", "lead_time_days"])
        suppliers_by_id = {s.id: s for s in self.store.list(Supplier, merchant_id)}

        def sort_key(offer: SupplierSku) -> tuple:
            keys: list[Any] = []
            for criterion in priority:
                if criterion == "preferred":
                    keys.append(0 if offer.preferred else 1)
                elif criterion == "cost":
                    keys.append(offer.cost)
                elif criterion == "lead_time_days" or criterion == "lead_time":
                    keys.append(offer.lead_time_days)
                elif criterion == "moq":
                    keys.append(offer.moq)
                elif criterion == "reliability":
                    supplier = suppliers_by_id.get(offer.supplier_id)
                    keys.append(-(supplier.reliability_score if supplier else 0.0))
            keys.append(offer.id)  # deterministic final tiebreaker
            return tuple(keys)

        return sorted(offers, key=sort_key)[0]

    # ================================================================================================
    # Procurement authority - deterministic merchant policy. AI never authorizes expenditure.
    # ================================================================================================

    def evaluate_po_authority(self, merchant_id: str, supplier_id: str, lines: list[dict[str, Any]]) -> tuple[Decision, dict[str, Any]]:
        config = self.store.get_config(merchant_id).get("procurement", {})
        spending = config.get("spending", {"auto_approve_limit": 0, "above_limit": "REQUIRE_APPROVAL"})
        supplier = self._supplier(merchant_id, supplier_id)
        reasoning: dict[str, Any] = {"checks": []}

        if supplier is None or supplier.status != "active":
            reasoning["checks"].append({"check": "supplier_active", "result": "FAIL", "detail": "unknown or inactive supplier"})
            return Decision.DENY, reasoning
        if any(line.get("unit_cost") is None or int(line.get("quantity_ordered", 0)) <= 0 for line in lines):
            reasoning["checks"].append({"check": "data_completeness", "result": "FAIL", "detail": "missing unit_cost or non-positive quantity"})
            return Decision.DENY, reasoning

        decision = Decision.ALLOW
        total_value = sum(int(line["quantity_ordered"]) * int(line["unit_cost"]) for line in lines)
        limit = int(spending.get("auto_approve_limit", 0))
        value_check = {"check": "spending_threshold", "total_value": total_value, "auto_approve_limit": limit}
        if total_value > limit:
            escalation = Decision(spending.get("above_limit", "REQUIRE_APPROVAL"))
            decision = _stricter(decision, escalation)
            value_check["result"] = escalation.value
        else:
            value_check["result"] = "PASS"
        reasoning["checks"].append(value_check)

        reliability_threshold = float(config.get("supplier_reliability_threshold", 0.0))
        reliability_check = {"check": "supplier_reliability", "score": supplier.reliability_score, "threshold": reliability_threshold}
        if supplier.reliability_score < reliability_threshold:
            decision = _stricter(decision, Decision.REQUIRE_APPROVAL)
            reliability_check["result"] = "REQUIRE_APPROVAL"
        else:
            reliability_check["result"] = "PASS"
        reasoning["checks"].append(reliability_check)

        cost_flags = []
        for line in lines:
            offer = self._offer(merchant_id, supplier_id, line["sku"])
            cost_decision, cost_info = self._cost_decision(merchant_id, offer.cost if offer else None, int(line["unit_cost"]))
            if cost_decision != Decision.ALLOW:
                decision = _stricter(decision, cost_decision)
                cost_flags.append({"sku": line["sku"], **cost_info, "result": cost_decision.value})
        reasoning["checks"].append({"check": "cost_variance", "flagged_lines": cost_flags, "result": "PASS" if not cost_flags else "ESCALATED"})

        unusual_flags = []
        for line in lines:
            unusual = self._unusual_quantity(merchant_id, supplier_id, line["sku"], int(line["quantity_ordered"]))
            multiplier_threshold = float(config.get("unusual_quantity_multiplier", 0) or 0)
            if unusual and multiplier_threshold and unusual["multiplier"] is not None and unusual["multiplier"] > multiplier_threshold:
                decision = _stricter(decision, Decision.REQUIRE_APPROVAL)
                unusual_flags.append({"sku": line["sku"], **unusual})
        reasoning["checks"].append({"check": "unusual_quantity", "flagged_lines": unusual_flags, "result": "PASS" if not unusual_flags else "REQUIRE_APPROVAL"})

        reasoning["final_decision"] = decision.value
        return decision, reasoning

    def _cost_decision(self, merchant_id: str, expected_cost: int | None, observed_cost: int) -> tuple[Decision, dict[str, Any]]:
        if expected_cost is None:
            # Phase 3 lesson applied: an unknown expected cost must NEVER silently become the observed
            # cost. There is nothing to compare against, so this can never resolve to ALLOW.
            return Decision.DENY, {"reason": "no_known_supplier_offer_cost", "observed_cost": observed_cost}
        config = self.store.get_config(merchant_id).get("procurement", {}).get("cost_tolerance", {})
        diff = observed_cost - expected_cost
        pct_variance = abs(diff) / expected_cost if expected_cost else (1.0 if diff else 0.0)
        tolerance_pct = float(config.get("pct", 0.0))
        tolerance_abs = int(config.get("absolute", 0))
        within = abs(diff) <= tolerance_abs or pct_variance <= tolerance_pct
        info = {
            "expected_cost": expected_cost, "observed_cost": observed_cost, "diff": diff,
            "pct_variance": round(pct_variance, 4), "tolerance_pct": tolerance_pct, "tolerance_abs": tolerance_abs,
            "within_tolerance": within,
        }
        if within:
            return Decision.ALLOW, info
        return Decision(config.get("above_tolerance", "REQUIRE_APPROVAL")), info

    def _unusual_quantity(self, merchant_id: str, supplier_id: str, sku: str, quantity: int) -> dict[str, Any] | None:
        pos_by_id = {p.id: p for p in self.store.list(PurchaseOrder, merchant_id) if p.supplier_id == supplier_id}
        history = [l.quantity_ordered for l in self.store.list(PurchaseOrderLine, merchant_id) if l.sku == sku and l.purchase_order_id in pos_by_id]
        if not history:
            return None
        average = sum(history) / len(history)
        return {"historical_average": round(average, 2), "requested": quantity, "multiplier": round(quantity / average, 2) if average else None}

    # ================================================================================================
    # PO lifecycle
    # ================================================================================================

    def create_purchase_order(self, merchant_id: str, supplier_id: str, lines: list[dict[str, Any]], *, idempotency_key: str | None = None) -> PurchaseOrder:
        """lines: [{"sku": ..., "quantity_ordered": ..., "location_ref": <optional>}] - cost/currency/MOQ
        are ALWAYS looked up from the on-file SupplierSku, never accepted from the caller. A line with no
        active SupplierSku is refused outright rather than invented.

        Step 5 (Part B/D/I): each line's DESTINATION location is decided HERE, exactly once, via Step 2's
        resolve_location_ref (explicit `location_ref` on the raw line wins; otherwise the merchant's
        configured default; otherwise the legacy "default" literal - full backward compatibility for
        callers that never pass location_ref at all). Once persisted on PurchaseOrderLine, this is the
        CANONICAL destination for the rest of this line's lifecycle - acknowledgement, shipment, and
        goods receipt all read it back rather than re-resolving. Different lines on the same PO may
        legitimately resolve to different locations (line-level, not header-level, matching the granularity
        ERPNext uses for its own per-line "Accepted Warehouse" - OSS finding, semantics only, Decision C)."""
        supplier = self._supplier(merchant_id, supplier_id)
        if supplier is None:
            raise ValueError(f"unknown supplier {supplier_id}")

        resolved_lines: list[dict[str, Any]] = []
        currency = None
        for raw in lines:
            offer = self._offer(merchant_id, supplier_id, raw["sku"])
            if offer is None:
                self.exceptions.create(
                    merchant_id=merchant_id, category="po_line_missing_supplier_offer",
                    message=f"No active SupplierSku for supplier={supplier_id} sku={raw['sku']} - refusing to invent cost/MOQ/lead time",
                    object_id=supplier_id, severity="error",
                )
                continue
            currency = currency or offer.currency
            location_ref = resolve_location_ref(self.store, merchant_id, explicit=raw.get("location_ref"))
            resolved_lines.append({
                "sku": raw["sku"], "supplier_sku": offer.supplier_sku, "quantity_ordered": int(raw["quantity_ordered"]),
                "unit_cost": offer.cost, "currency": offer.currency, "location_ref": location_ref,
            })

        po = PurchaseOrder(merchant_id=merchant_id, supplier_id=supplier_id, status="DRAFT" if resolved_lines else "BLOCKED", currency=currency or "INR", idempotency_key=idempotency_key)
        persisted_po = self.store.put(po)
        if persisted_po.id != po.id:
            return persisted_po  # idempotent replay of an already-created draft
        if not resolved_lines:
            self.audit.record(merchant_id=merchant_id, actor="procurement", source="procurement", action="purchase_order_blocked", object_type="PurchaseOrder", object_id=persisted_po.id, result="BLOCKED")
            return persisted_po
        for rl in resolved_lines:
            self.store.put(PurchaseOrderLine(
                merchant_id=merchant_id, purchase_order_id=persisted_po.id, sku=rl["sku"], supplier_sku=rl["supplier_sku"],
                location_ref=rl["location_ref"], quantity_ordered=rl["quantity_ordered"], unit_cost=rl["unit_cost"], currency=rl["currency"],
            ))
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="procurement", action="purchase_order_created", object_type="PurchaseOrder", object_id=persisted_po.id, result="DRAFT")
        return self.validate_purchase_order(merchant_id, persisted_po.id)

    def validate_purchase_order(self, merchant_id: str, po_id: str) -> PurchaseOrder:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status != "DRAFT":
            return po
        lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        decision, reasoning = self.evaluate_po_authority(merchant_id, po.supplier_id, [{"sku": l.sku, "quantity_ordered": l.quantity_ordered, "unit_cost": l.unit_cost} for l in lines])
        po.policy_decision = decision.value
        if decision == Decision.DENY:
            po.status = "BLOCKED"
            self.exceptions.create(merchant_id=merchant_id, category="po_denied", message=f"PO denied by policy: {reasoning}", object_id=po.id, severity="error")
        elif decision == Decision.REQUIRE_APPROVAL:
            po.status = "REQUIRES_APPROVAL"
            approval = Approval(merchant_id=merchant_id, action="purchase_order", object_id=po.id, requested_by="policy")
            self.store.put(approval)
            po.approval_id = approval.id
            self.exceptions.create(merchant_id=merchant_id, category="po_requires_approval", message=f"PO requires approval: {reasoning}", object_id=po.id, severity="warning")
        else:
            po.status = "AUTO_APPROVED"
        self.store.put(po)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="procurement", action="purchase_order_validated", object_type="PurchaseOrder", object_id=po.id, result=po.status)
        return po

    def approve_purchase_order(self, merchant_id: str, po_id: str, approver: str = "human") -> PurchaseOrder:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status != "REQUIRES_APPROVAL":
            return po
        po.status = "APPROVED"
        self.store.put(po)
        if po.approval_id:
            approval = self.store.get(Approval, merchant_id, po.approval_id)
            approval.status = "approved"
            approval.decided_by = approver
            approval.decided_at = now_utc()
            self.store.put(approval)
        # The "requires approval" notice this PO raised is now resolved by the human decision it was
        # flagging - it must not go on blocking PO closure forever as a permanently-open exception for
        # a review that has, in fact, already happened.
        for exc in self.store.list(ExceptionRecord, merchant_id):
            if exc.object_id == po.id and exc.category == "po_requires_approval" and exc.status == "open":
                exc.status = "resolved"
                self.store.put(exc)
        self.audit.record(merchant_id=merchant_id, actor=approver, source="procurement", action="purchase_order_approved", object_type="PurchaseOrder", object_id=po.id, result="APPROVED")
        return po

    def submit_purchase_order(self, merchant_id: str, po_id: str, *, simulate: str | None = None) -> PurchaseOrder:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status not in {"AUTO_APPROVED", "APPROVED"}:
            return po
        if self.supplier_connector is None:
            raise RuntimeError("supplier connector required for PO submission")
        lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        # Step 5B: `line_ref` is this PO's own canonical PurchaseOrderLine.id, sent as an opaque token -
        # the supplier is never asked to understand Sanocea's schema, only to echo it back exactly as
        # received on any later acknowledgement/shipment event (the smallest generic mechanism: no
        # separate reference-generation/mapping table, since po_line.id is already an opaque stable
        # string). This is what lets acknowledgement/inbound resolve to the EXACT line that was ordered,
        # even when two lines on the same PO share a SKU (see record_supplier_acknowledgement/
        # record_inbound_shipment's identity resolution).
        payload = {
            "purchase_order_id": po.id, "supplier_id": po.supplier_id, "currency": po.currency,
            "lines": [{"sku": l.sku, "supplier_sku": l.supplier_sku, "quantity_ordered": l.quantity_ordered, "unit_cost": l.unit_cost, "line_ref": l.id} for l in lines],
        } | ({"simulate": simulate} if simulate else {})
        idempotency_key = f"po-submit:{po.id}"
        try:
            result = self.supplier_connector.execute_mutation(
                MutationRequest(merchant_id=merchant_id, action="create_purchase_order", object_type="PurchaseOrder", payload=payload, idempotency_key=idempotency_key)
            )
        except TimeoutError:
            external = self.supplier_connector.find_purchase_order(merchant_id, po.id)
            if external:
                # The supplier DID create it before the response was lost - reconcile from their truth,
                # do not resubmit and do not fabricate a second external PO.
                po.external_ref = external["external_ref"]
                po.status = "SUBMITTED"
                po.submitted_at = now_utc()
                self.store.put(po)
                self.audit.record(merchant_id=merchant_id, actor="procurement", source=self.supplier_connector.name, action="purchase_order_submitted_recovered", object_type="PurchaseOrder", object_id=po.id, result="SUBMITTED")
                return po
            self.exceptions.create(merchant_id=merchant_id, category="po_submission_uncertain", message="PO submission mutation uncertain (timeout, supplier has no record of it) - NOT treated as success", object_id=po.id, severity="warning")
            self.audit.record(merchant_id=merchant_id, actor="procurement", source=self.supplier_connector.name, action="purchase_order_submission_uncertain", object_type="PurchaseOrder", object_id=po.id, result="uncertain")
            return po
        except RuntimeError as exc:
            category = "po_connector_rate_limited" if "429" in str(exc) else "po_connector_error"
            self.exceptions.create(merchant_id=merchant_id, category=category, message=str(exc), object_id=po.id, severity="warning")
            return po
        po.external_ref = result.external_ref
        po.status = "SUBMITTED"
        po.submitted_at = now_utc()
        self.store.put(po)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source=self.supplier_connector.name, action="purchase_order_submitted", object_type="PurchaseOrder", object_id=po.id, result="SUBMITTED")
        return po

    # ================================================================================================
    # Supplier acknowledgement - independent supplier-observed truth, cumulative, stale-event-safe.
    # ================================================================================================

    # CONFIRMED is deliberately included: a supplier can legitimately deliver a stale/out-of-order
    # acknowledgement event AFTER the PO has already reached full confirmation (e.g. a late-arriving
    # retry of an earlier partial ack). That must be evaluated by the sequence/staleness check below,
    # not rejected outright by a blanket PO-state guard.
    # OVER_CONFIRMED is included in both: a trailing ack (still evaluated for staleness) or a shipment
    # of the legitimately-ordered portion must remain possible after an over-confirmation is flagged -
    # the exception demands review, it does not have to halt the PO outright.
    _ACK_VALID_STATUSES = {"SUBMITTED", "PARTIALLY_CONFIRMED", "CONFIRMED", "OVER_CONFIRMED"}
    _SHIPMENT_VALID_STATUSES = {"CONFIRMED", "PARTIALLY_CONFIRMED", "OVER_CONFIRMED", "INBOUND"}
    _RECEIPT_VALID_STATUSES = {"INBOUND", "PARTIALLY_RECEIVED"}

    def record_supplier_acknowledgement(
        self, merchant_id: str, po_id: str, *, external_ref: str, sequence: int, lines: list[dict[str, Any]], status: str,
    ) -> SupplierAcknowledgement:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status not in self._ACK_VALID_STATUSES:
            # A DRAFT/BLOCKED/REQUIRES_APPROVAL/never-submitted PO cannot have been acknowledged by a
            # real supplier - refuse to let an event drive state that was never legitimately reachable.
            self.exceptions.create(
                merchant_id=merchant_id, category="acknowledgement_invalid_po_state",
                message=f"Acknowledgement received for PO {po_id} in state {po.status} (expected one of {sorted(self._ACK_VALID_STATUSES)}) - refused",
                object_id=po_id, severity="error",
            )
            ack = SupplierAcknowledgement(merchant_id=merchant_id, purchase_order_id=po_id, external_ref=external_ref, sequence=sequence, lines=lines, status=status, applied=False)
            return self.store.put(ack)
        # Step 7A.1: `applied` is no longer pre-computed via an unlocked read of "highest applied
        # sequence so far" here - that plain read raced against concurrent acknowledgements for the
        # SAME PO and could incorrectly mark a genuinely NEW, distinct, lower-numbered sequence as stale
        # merely because a numerically higher (but unrelated) sequence happened to be PROCESSED first
        # (found as a real, reproduced Postgres lost-update; see atomic_apply_acknowledgement's
        # docstring for the full root-cause/staleness-rule reasoning). The atomic per-line call below is
        # now the SOLE, race-free authority on staleness. `applied=False` here is only a placeholder to
        # satisfy the model while the DB-unique external_ref constraint does its own, unrelated,
        # unchanged duplicate-delivery dedup job.
        ack = SupplierAcknowledgement(merchant_id=merchant_id, purchase_order_id=po_id, external_ref=external_ref, sequence=sequence, lines=lines, status=status, applied=False)
        persisted = self.store.put(ack)
        if persisted.id != ack.id:
            return persisted  # duplicate external_ref - DB-enforced idempotency, no re-processing

        po_lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        any_over_confirmed = False
        event_applied: bool | None = None
        for line in lines:
            offer = self._offer(merchant_id, po.supplier_id, line["sku"])
            cost_decision, cost_info = self._cost_decision(merchant_id, offer.cost if offer else None, int(line["unit_cost"]))
            if cost_decision != Decision.ALLOW:
                self.exceptions.create(
                    merchant_id=merchant_id, category="po_cost_variance",
                    message=f"Acknowledged cost variance for {line['sku']} on PO {po_id}: {cost_info}",
                    object_id=po_id, severity="error" if cost_decision == Decision.DENY else "warning",
                )
            po_line, resolution_failure = self._resolve_po_line_for_event(po_lines, line_ref=line.get("line_ref"), sku=line["sku"])
            if resolution_failure:
                # Step 5B - a NEW failure mode: previously `_po_line`'s first-match lookup could never
                # detect ambiguity, it would silently credit whichever line happened to be listed first.
                # The pre-existing "SKU not found on this PO at all" case (resolution_failure is None,
                # po_line is None) is intentionally left as a silent no-op below, unchanged from before
                # Step 5B - only genuinely NEW ambiguity/bad-reference cases are flagged here.
                self.exceptions.create(
                    merchant_id=merchant_id, category=f"acknowledgement_{resolution_failure}",
                    message=(
                        f"Acknowledgement line for sku={line['sku']!r} line_ref={line.get('line_ref')!r} on PO {po_id} "
                        f"could not be resolved to exactly one PO line ({resolution_failure}) - refused, no mutation"
                    ),
                    object_id=po_id, severity="error",
                )
            if po_line is not None:
                # Step 7A.1 - ONE atomic transaction: the sequence-staleness claim, the PurchaseOrderLine
                # cumulative increment, and the location-scoped confirmed_inbound increment (capped to
                # the legitimate/ordered portion) all happen together, or none of them do. The TRUE
                # cumulative total is recorded on the line (visible/auditable even when it exceeds
                # quantity_ordered), but confirmed_inbound only ever receives the LEGITIMATE portion - an
                # over-confirming event can never push confirmed_inbound past what was actually ordered.
                raw_delta = int(line["quantity_confirmed"])
                result = self.store.atomic_apply_acknowledgement(merchant_id, po_id, po_line.id, line["sku"], po_line.location_ref, sequence, raw_delta)
                if event_applied is None:
                    event_applied = result["applied"]
                if not result["applied"]:
                    continue  # this exact sequence was already applied for this PO - handled once, below the loop
                if result["over_confirmed"]:
                    any_over_confirmed = True
                    self.exceptions.create(
                        merchant_id=merchant_id, category="supplier_overconfirmation",
                        message=(
                            f"Cumulative acknowledged quantity for {line['sku']} on PO {po_id} is {result['new_true_total']}, "
                            f"exceeding ordered quantity {po_line.quantity_ordered} - possible duplicate/erroneous "
                            f"remittance from supplier. Only {result['legitimate_delta']} of this event's {raw_delta} units "
                            f"were credited to confirmed_inbound; inventory position was NOT silently inflated."
                        ),
                        object_id=po_id, severity="error",
                    )

        if event_applied is False:
            # Step 7A.1 fix: `store.put()` cannot correct `applied` on an already-inserted
            # SupplierAcknowledgement row (PostgresStore's put() for this model is INSERT-only, ON
            # CONFLICT DO NOTHING, keyed on external_ref - a second put() call is silently discarded,
            # never updating existing fields; see mark_acknowledgement_applied's docstring for the full
            # story of this genuine, found bug). A direct, targeted update is required instead.
            self.store.mark_acknowledgement_applied(merchant_id, persisted.id, False)
            persisted.applied = False
            self.exceptions.create(
                merchant_id=merchant_id, category="stale_acknowledgement_ignored",
                message=f"Acknowledgement sequence {sequence} was already applied for PO {po_id} - evidence kept, state NOT altered",
                object_id=po_id, severity="warning",
            )
            self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier", action="acknowledgement_stale_ignored", object_type="PurchaseOrder", object_id=po_id, result="ignored")
            return persisted

        # event_applied is True, or None (no line resolved to a real PO line at all - nothing was
        # mutated, so there was nothing to atomically gate; treated as applied/evidence-only, matching
        # pre-7A.1 behavior for this edge case).
        self.store.mark_acknowledgement_applied(merchant_id, persisted.id, True)
        persisted.applied = True

        if status == "rejected":
            po.status = "REJECTED"
        elif any_over_confirmed:
            po.status = "OVER_CONFIRMED"
        else:
            po_lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
            fully_confirmed = bool(po_lines) and all(l.quantity_confirmed >= l.quantity_ordered for l in po_lines)
            po.status = "CONFIRMED" if fully_confirmed else "PARTIALLY_CONFIRMED"
        self.store.put(po)
        self._update_supplier_reliability(merchant_id, po.supplier_id)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier", action="acknowledgement_applied", object_type="PurchaseOrder", object_id=po_id, result=po.status)
        return persisted


    def _update_supplier_reliability(self, merchant_id: str, supplier_id: str) -> None:
        po_ids = {p.id for p in self.store.list(PurchaseOrder, merchant_id) if p.supplier_id == supplier_id and p.status not in {"DRAFT", "BLOCKED", "REQUIRES_APPROVAL", "APPROVED", "AUTO_APPROVED", "SUBMITTED"}}
        lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id in po_ids]
        if not lines:
            return
        ordered = sum(l.quantity_ordered for l in lines)
        confirmed = sum(min(l.quantity_confirmed, l.quantity_ordered) for l in lines)
        score = round(confirmed / ordered, 4) if ordered else 1.0
        supplier = self._supplier(merchant_id, supplier_id)
        if supplier:
            supplier.reliability_score = score
            self.store.put(supplier)

    # ================================================================================================
    # Inbound shipment (supplier's OWN manifest - not the warehouse's physical count)
    # ================================================================================================

    def record_inbound_shipment(
        self, merchant_id: str, po_id: str, *, external_shipment_ref: str, sequence: int, lines: list[dict[str, Any]], status: str = "dispatched",
    ) -> InboundShipment:
        """lines: [{"sku": ..., "quantity_shipped": int, "line_ref": <optional, PREFERRED - the exact
        canonical PurchaseOrderLine.id, echoed back exactly as sent on PO submission>}]

        Step 5B: resolved via the same exact-identity hierarchy as acknowledgement - `line_ref` first,
        else an unambiguous legacy SKU match on this PO. `quantity_shipped` accumulates on the resolved
        line only; a shipment for one line must never be attributed to a sibling line sharing its SKU.
        Preserves `PurchaseOrderLine.location_ref` unchanged (no location is read from or written by an
        inbound shipment event at all - it never mutates inventory - so there is no location to
        re-resolve or override here in the first place).
        """
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status not in self._SHIPMENT_VALID_STATUSES:
            self.exceptions.create(
                merchant_id=merchant_id, category="shipment_invalid_po_state",
                message=f"Shipment event received for PO {po_id} in state {po.status} (expected one of {sorted(self._SHIPMENT_VALID_STATUSES)}) - refused",
                object_id=po_id, severity="error",
            )
            shipment = InboundShipment(merchant_id=merchant_id, purchase_order_id=po_id, supplier_id=po.supplier_id, external_shipment_ref=external_shipment_ref, sequence=sequence, status=status, applied=False)
            return self.store.put(shipment)
        highest_applied = max(
            (s.sequence for s in self.store.list(InboundShipment, merchant_id) if s.purchase_order_id == po_id and s.applied),
            default=-1,
        )
        applied = sequence > highest_applied
        shipment = InboundShipment(merchant_id=merchant_id, purchase_order_id=po_id, supplier_id=po.supplier_id, external_shipment_ref=external_shipment_ref, sequence=sequence, status=status, applied=applied, dispatched_at=now_utc())
        persisted = self.store.put(shipment)
        if persisted.id != shipment.id:
            return persisted  # duplicate external_shipment_ref - DB-enforced idempotency

        if not applied:
            self.exceptions.create(merchant_id=merchant_id, category="stale_shipment_event_ignored", message=f"Shipment sequence {sequence} <= highest applied {highest_applied} for PO {po_id} - evidence kept, state NOT altered", object_id=po_id, severity="warning")
            self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier", action="shipment_stale_ignored", object_type="PurchaseOrder", object_id=po_id, result="ignored")
            return persisted

        po_lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        for line in lines:
            po_line, resolution_failure = self._resolve_po_line_for_event(po_lines, line_ref=line.get("line_ref"), sku=line["sku"])
            if resolution_failure:
                self.exceptions.create(
                    merchant_id=merchant_id, category=f"shipment_{resolution_failure}",
                    message=(
                        f"Shipment line for sku={line['sku']!r} line_ref={line.get('line_ref')!r} on PO {po_id} could "
                        f"not be resolved to exactly one PO line ({resolution_failure}) - refused, no mutation"
                    ),
                    object_id=po_id, severity="error",
                )
            self.store.put(InboundShipmentLine(
                merchant_id=merchant_id, inbound_shipment_id=persisted.id, sku=line["sku"],
                quantity_shipped=int(line["quantity_shipped"]), purchase_order_line_id=po_line.id if po_line else None,
            ))
            if po_line is not None:
                po_line.quantity_shipped += int(line["quantity_shipped"])
                self.store.put(po_line)
        if po.status in {"CONFIRMED", "PARTIALLY_CONFIRMED"}:
            po.status = "INBOUND"
            self.store.put(po)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="supplier", action="shipment_recorded", object_type="PurchaseOrder", object_id=po_id, result=persisted.status)
        return persisted

    # ================================================================================================
    # Goods receipt - the ONLY point where confirmed inbound becomes sellable inventory, and only for
    # lines Sanocea can confidently resolve. Uncertain receipt never silently touches sellable stock.
    # ================================================================================================

    def record_goods_receipt(
        self, merchant_id: str, po_id: str, *, external_receipt_ref: str, lines: list[dict[str, Any]], inbound_shipment_id: str | None = None,
    ) -> GoodsReceipt:
        """lines: [{"raw_sku_reference": <what is physically labeled>, "quantity_received": int,
        "po_line_id": <optional, PREFERRED - the exact canonical PurchaseOrderLine.id>,
        "location_ref": <optional - what the receiving payload CLAIMS the receiving location is>}]

        Step 5A (PO-LINE RECEIPT IDENTITY): a receipt line must resolve to EXACTLY ONE canonical
        PurchaseOrderLine before any inventory mutation, via a strict hierarchy - never SKU alone as
        authoritative identity, and never list/dict insertion order to break a tie:

          1. `po_line_id`, if supplied - the strongest identity available (this PO's own canonical line
             id). Must belong to THIS PO or the line is refused as `unknown_po_line`; no fallback to SKU
             matching is attempted once an explicit id is given.
          2. Otherwise, `raw_sku_reference` matched against this PO's own lines (legacy input, backward
             compatible with pre-Step-5A payloads). Resolves ONLY when EXACTLY ONE PurchaseOrderLine on
             this PO carries that SKU. Two or more matches - the same SKU ordered to two different
             locations, or twice to the SAME location (a strictly harder case: location cannot
             disambiguate it either) - is `ambiguous_po_line`: refused outright, no mutation, no
             guessing. (sku, location_ref) is deliberately NOT used as a disambiguating composite key -
             the canonical model does not guarantee that pair is unique, so treating it as an identity
             would be exactly the "weaker, must fail closed" tier this hierarchy exists to avoid leaning
             on when a stronger identity (po_line_id) is available instead.

        Step 5 (Part F/H, unchanged): physical receipt always mutates inventory at the RESOLVED line's
        OWN canonical `location_ref` - never a location named in the payload. `location_ref` on an
        incoming line is VALIDATION EVIDENCE ONLY: it can flag a mismatch against the resolved line's
        canonical destination, but it never selects which line resolves, and never chooses which
        location is mutated - no override path exists.
        """
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status not in self._RECEIPT_VALID_STATUSES:
            self.exceptions.create(
                merchant_id=merchant_id, category="goods_receipt_invalid_po_state",
                message=f"Goods receipt recorded for PO {po_id} in state {po.status} (expected one of {sorted(self._RECEIPT_VALID_STATUSES)}) - refused, inventory NOT touched",
                object_id=po_id, severity="error",
            )
            receipt = GoodsReceipt(merchant_id=merchant_id, purchase_order_id=po_id, inbound_shipment_id=inbound_shipment_id, external_receipt_ref=external_receipt_ref, status="received_with_exceptions")
            return self.store.put(receipt)
        po_lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        po_lines_by_id = {l.id: l for l in po_lines}
        po_lines_by_sku: dict[str, list[PurchaseOrderLine]] = {}
        for l in po_lines:
            po_lines_by_sku.setdefault(l.sku, []).append(l)
        receipt = GoodsReceipt(merchant_id=merchant_id, purchase_order_id=po_id, inbound_shipment_id=inbound_shipment_id, external_receipt_ref=external_receipt_ref)
        persisted_receipt = self.store.put(receipt)
        if persisted_receipt.id != receipt.id:
            return persisted_receipt  # duplicate external_receipt_ref - DB-enforced idempotency

        has_exception = False
        for raw in lines:
            sku_ref = raw.get("raw_sku_reference")
            quantity_received = int(raw["quantity_received"])
            explicit_po_line_id = raw.get("po_line_id")

            po_line: PurchaseOrderLine | None = None
            disposition: str | None = None
            failure_message: str | None = None
            if explicit_po_line_id:
                po_line = po_lines_by_id.get(explicit_po_line_id)
                if po_line is None:
                    disposition = "unknown_po_line"
                    failure_message = f"Receipt referenced po_line_id={explicit_po_line_id!r} which does not belong to PO {po.id} - refused, inventory NOT touched"
            else:
                candidates = po_lines_by_sku.get(sku_ref, []) if sku_ref else []
                if len(candidates) == 1:
                    po_line = candidates[0]
                elif len(candidates) > 1:
                    disposition = "ambiguous_po_line"
                    failure_message = (
                        f"Receipt reference {sku_ref!r} matches {len(candidates)} PO lines on PO {po.id} "
                        f"(same SKU on multiple lines/locations) - ambiguous without an explicit po_line_id, "
                        f"refused rather than guessed, inventory NOT touched"
                    )
                else:
                    disposition = "wrong_sku" if sku_ref else "unexpected_item"
                    failure_message = f"Received {quantity_received} units of unrecognized reference {sku_ref!r} against PO {po.id} - missing/wrong PO reference"

            if po_line is None:
                self.store.put(GoodsReceiptLine(
                    merchant_id=merchant_id, goods_receipt_id=persisted_receipt.id, sku=None, raw_sku_reference=sku_ref,
                    quantity_received=quantity_received, quantity_expected=None, disposition=disposition,
                    location_ref=None, purchase_order_line_id=None,
                ))
                self.exceptions.create(merchant_id=merchant_id, category=f"goods_receipt_{disposition}", message=failure_message, object_id=po.id, severity="error")
                has_exception = True
                continue  # Uncertain/ambiguous receipt: do NOT touch inventory for an unresolved reference.

            claimed_location = raw.get("location_ref")
            if claimed_location and claimed_location != po_line.location_ref:
                # Part H - a wrong-location receipt must be FLAGGED, never silently mutate the location
                # the payload claims. The canonical PO-line destination remains authoritative below
                # regardless; no override path exists for Step 5.
                self.exceptions.create(
                    merchant_id=merchant_id, category="goods_receipt_wrong_location_attempted",
                    message=(
                        f"Receipt for {po_line.sku} (po_line_id={po_line.id}) on PO {po.id} claimed location "
                        f"{claimed_location!r}, but this PO line's canonical destination is {po_line.location_ref!r} "
                        f"- receiving at the claimed location was refused; inventory was applied at the canonical "
                        f"destination only."
                    ),
                    object_id=po.id, severity="error",
                )
                has_exception = True

            expected_remaining = po_line.quantity_ordered - po_line.quantity_received
            if quantity_received < expected_remaining:
                disposition = "shortage"
            elif quantity_received > expected_remaining:
                disposition = "excess"
            else:
                disposition = "match"
            self.store.put(GoodsReceiptLine(
                merchant_id=merchant_id, goods_receipt_id=persisted_receipt.id, sku=po_line.sku, raw_sku_reference=sku_ref,
                quantity_received=quantity_received, quantity_expected=expected_remaining, disposition=disposition,
                location_ref=po_line.location_ref, purchase_order_line_id=po_line.id,
            ))
            if disposition != "match":
                self.exceptions.create(
                    merchant_id=merchant_id, category=f"goods_receipt_{disposition}",
                    message=f"{disposition} for {po_line.sku} (po_line_id={po_line.id}) on PO {po.id}: expected {expected_remaining}, received {quantity_received}",
                    object_id=po.id, severity="warning",
                )
                has_exception = True

            # The physical count itself is real regardless of the discrepancy classification - the
            # ACTUAL counted quantity (never the expected/claimed quantity) is what becomes sellable, and
            # it is applied at the EXACT RESOLVED LINE's own canonical location only (Part F/Step 5A) -
            # never the payload's claimed location (Part H), and never a different line sharing the
            # same SKU.
            po_line.quantity_received += quantity_received
            self.store.put(po_line)
            self._apply_goods_receipt_to_inventory(merchant_id, po_line.sku, po_line.location_ref, quantity_received)

        persisted_receipt = self.store.get(GoodsReceipt, merchant_id, persisted_receipt.id)
        persisted_receipt.status = "received_with_exceptions" if has_exception else "received"
        self.store.put(persisted_receipt)
        self._advance_po_receipt_status(merchant_id, po_id)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="warehouse", action="goods_receipt_recorded", object_type="PurchaseOrder", object_id=po_id, result=persisted_receipt.status)
        return persisted_receipt

    def _apply_goods_receipt_to_inventory(self, merchant_id: str, sku: str, location_ref: str, quantity_received: int) -> None:
        # Step 5 (Part F): increase quantity/available and reduce confirmed_inbound at the PO line's
        # OWN destination location only - every other location's Inventory row is untouched by this call.
        self.store.atomic_adjust_inventory(
            merchant_id, sku, location_ref,
            confirmed_inbound_delta=-quantity_received, quantity_delta=quantity_received, available_delta=quantity_received,
        )

    def _advance_po_receipt_status(self, merchant_id: str, po_id: str) -> None:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        lines = [l for l in self.store.list(PurchaseOrderLine, merchant_id) if l.purchase_order_id == po_id]
        if not lines:
            return
        fully_received = all(l.quantity_received >= l.quantity_ordered for l in lines)
        any_received = any(l.quantity_received > 0 for l in lines)
        if fully_received:
            po.status = "RECEIVED"
        elif any_received:
            po.status = "PARTIALLY_RECEIVED"
        self.store.put(po)

    def reconcile_purchase_order(self, merchant_id: str, po_id: str) -> PurchaseOrder:
        po = self.store.get(PurchaseOrder, merchant_id, po_id)
        if po.status not in {"RECEIVED", "PARTIALLY_RECEIVED"}:
            return po
        open_exceptions = [e for e in self.store.list(ExceptionRecord, merchant_id) if e.object_id == po.id and e.status == "open"]
        if open_exceptions:
            return po
        po.status = "RECONCILED"
        self.store.put(po)
        po.status = "CLOSED"
        self.store.put(po)
        self.audit.record(merchant_id=merchant_id, actor="procurement", source="procurement", action="purchase_order_closed", object_type="PurchaseOrder", object_id=po.id, result="CLOSED")
        return po

    # ================================================================================================
    # Lookups
    # ================================================================================================

    def _supplier(self, merchant_id: str, supplier_id: str) -> Supplier | None:
        return next((s for s in self.store.list(Supplier, merchant_id) if s.id == supplier_id), None)

    def _offer(self, merchant_id: str, supplier_id: str, sku: str) -> SupplierSku | None:
        return next((o for o in self.store.list(SupplierSku, merchant_id) if o.supplier_id == supplier_id and o.sku == sku and o.status == "active"), None)

    def _resolve_po_line_for_event(
        self, po_lines: list[PurchaseOrderLine], *, line_ref: str | None, sku: str | None,
    ) -> tuple[PurchaseOrderLine | None, str | None]:
        """Step 5B - shared exact PO-line identity resolution for supplier acknowledgement and inbound
        shipment events (goods receipt has its own, separately-accepted Step 5A resolution with
        different zero-match handling - deliberately kept independent, not unified with this one).
        Never SKU alone, supplier SKU alone, or (sku, location_ref) as authoritative identity - the
        canonical model does not guarantee (sku, location_ref) is unique, so it is never used to select
        between candidates anywhere in this hierarchy.

        `po_lines` is this PO's own lines, pre-fetched once by the caller (avoids re-querying the store
        once per line in a multi-line event).

        Returns (po_line, disposition):
          - (line, None) - resolved to exactly one line via `line_ref` (preferred - this PO's own
            canonical PurchaseOrderLine.id, echoed back exactly as sent on PO submission) or an
            unambiguous legacy SKU match.
          - (None, "unknown_po_line") - an explicit `line_ref` was supplied but does not belong to this
            PO; no fallback to SKU matching is attempted once an explicit reference is given.
          - (None, "ambiguous_po_line") - no `line_ref` given, and more than one line on this PO carries
            the supplied SKU (the same SKU ordered to two different locations, or twice to the SAME
            location) - refused rather than guessed; never resolved by list/dict insertion order.
          - (None, None) - no `line_ref` given and no line on this PO carries the supplied SKU at all.
            This is the PRE-EXISTING "unknown SKU" case (unchanged since before Step 5B) and is
            intentionally left silent here (no exception, matching the caller's prior behavior) - only
            the genuinely NEW ambiguous/bad-reference cases are surfaced as exceptions by the caller.
        """
        if line_ref:
            po_line = next((l for l in po_lines if l.id == line_ref), None)
            return (po_line, None) if po_line is not None else (None, "unknown_po_line")
        candidates = [l for l in po_lines if l.sku == sku] if sku else []
        if len(candidates) == 1:
            return candidates[0], None
        if len(candidates) > 1:
            return None, "ambiguous_po_line"
        return None, None
